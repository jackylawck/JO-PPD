#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Jumbo Orient Punch Problem Detector (東淦打卡異常偵測工具)
- 本機離線運行，保障員工考勤數據私隱
- 動態表頭識別與欄位保護
- 向量化執行審計規則，生成原廠對齊格式之 Problem.xlsx
"""

import os
import re
import glob
import yaml
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CONFIG_PATH = "rules.yaml"

def locate_attendance_files():
    """搜尋目錄下最新考勤原檔，排除暫存與輸出檔"""
    files = glob.glob("*Attendance*.xls*") + glob.glob("*attendance*.xls*")
    files = [f for f in files if "problem" not in f.lower() and not os.path.basename(f).startswith("~$")]
    if not files:
        files = [f for f in glob.glob("*.xls*") if "problem" not in f.lower() and not os.path.basename(f).startswith("~$")]
    if not files:
        raise FileNotFoundError("目錄下找不到考勤原始 Excel 檔案！")
    
    files.sort(key=lambda x: os.path.getmtime(x), reverse=True)
    return files[0]

def parse_excel_structure(file_path, cutoff_day_cfg=25):
    """自適應引擎、動態定位表頭、保護欄位長度並抽取 Period"""
    engine = "openpyxl" if file_path.endswith(".xlsx") else "xlrd"
    df_raw = pd.read_excel(file_path, header=None, engine=engine)

    # 1. 抽取 Period 檔名與截數截止日
    out_filename = "Problem.xlsx"
    cutoff_date = None
    period_pattern = re.compile(r'Period\s*:\s*:?\s*(\d{1,2})[-/](\d{1,2})[-/](\d{4})\s*To\s*(\d{1,2})[-/](\d{1,2})[-/](\d{4})', re.IGNORECASE)
    
    for r in range(min(10, len(df_raw))):
        line = " ".join(df_raw.iloc[r].dropna().astype(str).values)
        m = period_pattern.search(line)
        if m:
            d_start, m_start, y_start, d_end, m_end, y_end = m.groups()
            out_filename = f"{int(d_start)}-{int(d_end)}.{int(m_end)}.{y_end}.Problem.xlsx"
            if cutoff_day_cfg:
                cutoff_date = pd.Timestamp(year=int(y_end), month=int(m_end), day=int(cutoff_day_cfg))
            break

    # 2. 定位表頭
    header_idx = None
    for idx, row in df_raw.iterrows():
        if any("Department Code" in str(v) for v in row.values):
            header_idx = idx
            break

    if header_idx is None:
        raise ValueError("無法在原始檔案中定位 'Department Code' 表頭！")

    headers = [
        'Dept_Code', 'Dept_Name', 'Emp_Code', 'Name', 'Designation',
        'Date', 'Day', 'Clk1', 'Clk2', 'Clk3', 'Clk4', 'Time_In', 'Time_Out',
        'Shift', 'Normal', 'Late', 'Early_Out', 'Actual',
        'OT_1_0', 'OT_1_5', 'OT_2_0', 'OT_3_0', 'Flat', 'Reason'
    ]

    # 3. 欄位數安全長度防護
    data_start = header_idx + 2
    df_data = df_raw.iloc[data_start:, :].copy()
    if df_data.shape[1] < len(headers):
        for i in range(len(headers) - df_data.shape[1]):
            df_data[df_data.shape[1]] = 0.0
    
    df_data = df_data.iloc[:, :len(headers)]
    df_data.columns = headers
    return df_data, out_filename, cutoff_date

def load_rules(path=CONFIG_PATH):
    if not os.path.exists(path):
        raise FileNotFoundError(f"找不到配置檔: {path}")
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)

def run_integrity_audit(df, config, cutoff_date=None):
    original_dept_order = list(dict.fromkeys(df['Dept_Code'].dropna()))
    
    df = df[df['Emp_Code'].notna()].copy()
    df['Emp_Code'] = df['Emp_Code'].astype(str).str.strip().str.upper()
    
    # 僅審計 E-code
    df = df[df['Emp_Code'].str.startswith('E')].copy()
    
    # 部門白名單過濾
    target_depts = config.get('target_departments', [])
    if target_depts:
        df = df[df['Dept_Code'].isin(target_depts)].copy()

    df['Date'] = pd.to_datetime(df['Date'], errors='coerce')

    # 截數日過濾
    if cutoff_date is not None:
        df = df[df['Date'] <= cutoff_date].copy()

    num_cols = ['Clk1', 'Clk2', 'Clk3', 'Clk4', 'Time_In', 'Time_Out', 'Normal', 'Late', 'Early_Out', 'Actual']
    for col in num_cols:
        df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0.0)

    for c in ['Dept_Code', 'Dept_Name', 'Name', 'Designation', 'Shift', 'Reason']:
        df[c] = df[c].fillna('').astype(str).str.strip().str.upper()

    # 向量化求取最後打卡時間
    df['Last_Clk'] = df[['Clk1', 'Clk2', 'Clk3', 'Clk4']].max(axis=1)

    # 向量化比對規則
    anomaly_mask = pd.Series(False, index=df.index)
    for r in config.get('rules', []):
        cond = r.get('condition')
        try:
            eval_res = df.eval(cond)
            anomaly_mask |= eval_res
        except Exception as e:
            print(f"[!] 規則解析警告 ({r.get('id')}): {e}")

    problem_df = df[anomaly_mask].copy()

    # 原廠格式：當打卡記錄全為 0 時，原廠輸出 Time IN/OUT 強制歸零
    zero_clk_mask = (problem_df['Clk1'] == 0) & (problem_df['Clk2'] == 0) & (problem_df['Clk3'] == 0) & (problem_df['Clk4'] == 0)
    problem_df.loc[zero_clk_mask, 'Time_In'] = 0.0
    problem_df.loc[zero_clk_mask, 'Time_Out'] = 0.0

    # 保持原廠部門出現排版
    dept_cat = pd.CategoricalDtype(categories=original_dept_order, ordered=True)
    problem_df['Dept_Order'] = problem_df['Dept_Code'].astype(dept_cat)
    problem_df = problem_df.sort_values(by=['Dept_Order', 'Emp_Code', 'Date'])

    return problem_df

def export_aligned_excel(df_problems, output_filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = output_filename.replace(".xlsx", ".Log")[:31]

    # 雙層表頭
    ws.append(['Department Code', 'Department Name', 'Emp Code', 'Name', 'Designation',
               'Date', 'Day', 'Actual Clocking', None, None, None, 'Time', None])
    ws.append([None, None, None, None, None, None, None, None, None, None, None, 'IN', 'OUT'])

    ws.merge_cells('H1:K1')
    ws.merge_cells('L1:M1')

    cols_to_export = ['Dept_Code', 'Dept_Name', 'Emp_Code', 'Name', 'Designation',
                      'Date', 'Day', 'Clk1', 'Clk2', 'Clk3', 'Clk4', 'Time_In', 'Time_Out']
    
    rows_data = df_problems[cols_to_export].values.tolist()
    for row in rows_data:
        if pd.notna(row[5]):
            row[5] = pd.to_datetime(row[5]).to_pydatetime()
        ws.append(row)

    # 原廠樣式標準：Verdana 字型、灰細點邊框
    header_font = Font(name="Verdana", size=10, bold=True)
    body_font = Font(name="Verdana", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    dotted_border = Border(
        left=Side(style='dotted', color='A0A0A0'),
        right=Side(style='dotted', color='A0A0A0'),
        top=Side(style='thin', color='A0A0A0'),
        bottom=Side(style='thin', color='A0A0A0')
    )

    for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=13):
        for cell in row:
            cell.font = header_font
            cell.alignment = center_align

    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=13):
        for col_idx, cell in enumerate(row, start=1):
            cell.font = body_font
            cell.border = dotted_border
            if col_idx == 6:
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
                cell.alignment = center_align
            elif isinstance(cell.value, (int, float)):
                cell.number_format = '0.00'

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col[:20])
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_filename)
    print(f"[✓] 成功產出 100% 對齊報表: {output_filename} (共篩選出 {len(df_problems)} 筆異常記錄)")

def main():
    try:
        config = load_rules()
        file_path = locate_attendance_files()
        print(f"[*] 分析考勤原檔: {file_path}")
        
        cutoff_day_cfg = config.get('cutoff_day', 25)
        df_data, out_name, cutoff_date = parse_excel_structure(file_path, cutoff_day_cfg)
        
        if cutoff_date:
            print(f"[*] 生效結算截止日 (Cut-off): <= {cutoff_date.strftime('%Y-%m-%d')}")
            
        problems = run_integrity_audit(df_data, config, cutoff_date)
        export_aligned_excel(problems, out_name)
    except Exception as e:
        print(f"[X] 執行中斷: {e}")

if __name__ == '__main__':
    main()
